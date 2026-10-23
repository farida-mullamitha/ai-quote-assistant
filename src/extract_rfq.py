"""Read an RFQ file (PDF, Word or Excel) and return its line items.

    from extract_rfq import extract_rfq_lines
    lines = extract_rfq_lines("data/raw/RFQ_KBS-0147.pdf")

or from the command line:
    python src/extract_rfq.py data/raw/RFQ_KBS-0147.pdf

How it works (same idea for all three formats)
----------------------------------------------
1. Pull every TABLE out of the file as a plain list of rows.
2. In each table, find the HEADER row by looking for column names we know
   ("Description", "Part Code", "Qty", ...). Different customers name the
   columns differently, so each canonical field has several synonyms.
3. Every row under the header is one RFQ line. Clean the text and return
   one table with the same columns, whatever the file type was.
"""

import re
import sys
from pathlib import Path

import pandas as pd

# ------------------------------------------------------------ column names
# Canonical field -> the header texts customers might use for it.
# Headers are compared after lower-casing and removing everything except
# letters and digits, so "Part Code", "part-code" and "PART CODE" all match.
SYNONYMS = {
    "line_no": {"no", "sno", "srno", "slno", "item", "itemno", "sr", "serial", "pos", "position"},
    "description": {"description", "itemdescription", "details", "particulars", "materialdescription", "product"},
    "make": {"make", "brand", "manufacturer", "supplier", "mfr"},
    "part_code": {"partcode", "partno", "partnumber", "catno", "catalogueno", "catalognumber",
                  "ordercode", "orderingcode", "code", "reference", "modelno", "mlfb"},
    "qty": {"qty", "quantity", "qnty"},
    "unit": {"unit", "uom", "units", "unitofmeasure"},
}

OUTPUT_COLUMNS = [
    "line_no", "description", "make", "part_code", "code_note",
    "qty", "unit", "source",
]

# Words people write instead of a part code when they have none.
NO_CODE_WORDS = {"", "TBC", "TBA", "NA", "NIL", "NONE", "TOBECONFIRMED"}


# ------------------------------------------------------------ small helpers
def clean(cell) -> str:
    """Make a cell into tidy one-line text."""
    if cell is None:
        return ""
    text = str(cell).replace("-\n", "-")          # 'thermal-\nmagnetic'
    return re.sub(r"\s+", " ", text).strip()


def header_key(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", clean(text).lower())


def map_header(row):
    """If this row looks like a header, return {column_index: field}."""
    mapping = {}
    for index, cell in enumerate(row):
        key = header_key(cell)
        for field, names in SYNONYMS.items():
            if key in names and field not in mapping.values():
                mapping[index] = field
                break
    # A real header names the description plus at least two other fields.
    if "description" in mapping.values() and len(mapping) >= 3:
        return mapping
    return None


def split_code_note(cell):
    """Separate a real part code from explanatory notes."""
    text = clean(cell)

    # Extract text written inside brackets.
    notes = re.findall(r"\(([^)]*)\)", text)

    # Remove bracketed notes from the possible part code.
    code = re.sub(r"\([^)]*\)", "", text).strip()

    # Compare a simplified version of the text.
    key = header_key(code)

    # Recognise common ways of saying a code is missing.
    missing_code = (
        key in {"", "tbc", "tba", "na", "nil", "none", "tobeconfirmed"}
        or "tobeconfirmed" in key
        or (
            "supplier" in key
            and ("code" in key or "particular" in key)
        )
    )

    if missing_code:
        if text and not notes:
            notes.append(text)
        code = ""

    return code, "; ".join(note.strip() for note in notes)



def normalise_make(text: str) -> str:
    """Use the same supplier names as the price catalogue."""
    low = clean(text).lower()
    if "abb" in low:
        return "ABB"
    if "schneider" in low:
        return "Schneider Electric"
    return clean(text)


def to_number(text: str):
    match = re.search(r"\d+(?:\.\d+)?", clean(text).replace(",", ""))
    if not match:
        return None
    number = float(match.group(0))
    return int(number) if number.is_integer() else number


# ------------------------------------------------------------ table -> lines
def lines_from_tables(tables):
    """tables: list of (location_text, rows). Returns a DataFrame of lines."""
    found = []
    last_mapping = None
    last_width = None

    for location, rows in tables:
        rows = [list(r) for r in rows if any(clean(c) for c in r)]
        if not rows:
            continue

        # Find the header row inside this table (could be below a title).
        mapping, start = None, 0
        for i, row in enumerate(rows[:30]):
            mapping = map_header(row)
            if mapping:
                start = i + 1
                break

        # A table that continues on the next page may not repeat its header.
        if mapping is None and last_mapping and len(rows[0]) == last_width:
            mapping, start = last_mapping, 0
        if mapping is None:
            continue
        last_mapping, last_width = mapping, len(rows[0])

        for row in rows[start:]:
            if map_header(row):          # repeated header row: skip
                continue
            values = {field: clean(row[i]) for i, field in mapping.items() if i < len(row)}
            # The part code cell needs special care (notes in brackets).
            code_cells = [row[i] for i, f in mapping.items() if f == "part_code" and i < len(row)]
            code, note = split_code_note(code_cells[0] if code_cells else "")

            line_no = values.get("line_no", "")
            description = values.get("description", "")
            qty = to_number(values.get("qty", ""))
            # Skip total / blank / note rows: a real line has a description
            # (or code) AND a line number or quantity.
            if not (description or code) or not (line_no or qty is not None):
                continue

            found.append(
                {
                    "line_no": line_no,
                    "description": description,
                    "make": normalise_make(values.get("make", "")),
                    "part_code": code,
                    "code_note": note,
                    "qty": qty,
                    "unit": values.get("unit", ""),
                    "source": location,
                }
            )

    return pd.DataFrame(found, columns=OUTPUT_COLUMNS)


# ------------------------------------------------------------ readers
def tables_from_pdf(path: Path):
    import pdfplumber

    tables = []
    with pdfplumber.open(path) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            for t, rows in enumerate(page.extract_tables(), start=1):
                tables.append((f"page {number}, table {t}", rows))
    return tables


def tables_from_docx(path: Path):
    import docx

    document = docx.Document(str(path))
    tables = []
    for t, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        tables.append((f"table {t}", rows))
    return tables


def tables_from_excel(path: Path):
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    tables = []
    for sheet in workbook.worksheets:
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        tables.append((f"sheet '{sheet.title}'", rows))
    return tables


READERS = {
    ".pdf": tables_from_pdf,
    ".docx": tables_from_docx,
    ".xlsx": tables_from_excel,
    ".xlsm": tables_from_excel,
}


def extract_rfq_lines(path) -> pd.DataFrame:
    """Return the RFQ's line items as a table (one row per RFQ line)."""
    path = Path(path)
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise ValueError(
            f"Unsupported file type '{path.suffix}'. Use PDF, DOCX or XLSX."
        )
    lines = lines_from_tables(reader(path))
    if lines.empty:
        raise ValueError(
            f"No line-item table found in {path.name}. If this is a scanned "
            "PDF it needs OCR first; if the items are in an email or plain "
            "text, they need the AI extraction step (not built yet)."
        )
    return lines


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Usage: python src/extract_rfq.py <rfq file>")
    result = extract_rfq_lines(sys.argv[1])
    pd.set_option("display.width", 200, "display.max_colwidth", 38)
    print(result.drop(columns=["source"]).to_string(index=False))
    print(f"\n{len(result)} lines extracted.")