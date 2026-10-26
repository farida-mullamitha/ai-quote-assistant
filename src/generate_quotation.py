from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import range_boundaries

from src.extract_rfq import extract_rfq_lines
from src.load_price_lists import load_catalogue, add_supplier_prices


PROJECT_ROOT = Path(__file__).resolve().parent.parent

TEMPLATE_PATH = PROJECT_ROOT / "data" / "raw" / "IEPCI Quotation .xlsx"
RFQ_PATH = PROJECT_ROOT / "data" / "raw" / "RFQ_KBS-0147.pdf"
OUTPUT_DIR = PROJECT_ROOT / "data" / "processed"

FIRST_ITEM_ROW = 13
LAST_ITEM_ROW = 29

# Template column positions
COL_NUMBER = 1
COL_PART_CODE = 2
COL_DESCRIPTION = 3
COL_UNIT = 9
COL_QUANTITY = 10
COL_UNIT_PRICE = 11
COL_LINE_TOTAL = 13


def get_value(item, *column_names, default=None):
    """Return the first available, non-empty value from an RFQ item."""
    for name in column_names:
        if name in item.index:
            value = item[name]

            if value is not None and not pd.isna(value):
                if str(value).strip():
                    return value

    return default


def copy_row_format(sheet, source_row, target_row):
    """Copy formatting and row height without copying cell values."""
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)

        if isinstance(source, MergedCell):
            continue

        if source.has_style:
            target._style = copy(source._style)

        if source.number_format:
            target.number_format = source.number_format

        if source.alignment:
            target.alignment = copy(source.alignment)

        if source.protection:
            target.protection = copy(source.protection)

    source_height = sheet.row_dimensions[source_row].height
    if source_height is not None:
        sheet.row_dimensions[target_row].height = source_height


def expand_item_table(sheet, item_count):
    """Expand the item table while preserving existing merged cells."""
    current_capacity = LAST_ITEM_ROW - FIRST_ITEM_ROW + 1
    extra_rows = item_count - current_capacity

    if extra_rows <= 0:
        return

    insert_at = LAST_ITEM_ROW + 1

    # Save and temporarily remove merged ranges below the item table.
    affected_ranges = [
        str(merged_range)
        for merged_range in list(sheet.merged_cells.ranges)
        if merged_range.max_row >= insert_at
    ]

    for merged_range in affected_ranges:
        sheet.unmerge_cells(merged_range)

    # Insert rows for additional RFQ items.
    sheet.insert_rows(insert_at, extra_rows)

    # Restore existing merged ranges at their new row positions.
    for merged_range in affected_ranges:
        min_col, min_row, max_col, max_row = range_boundaries(
            merged_range
        )

        if min_row >= insert_at:
            min_row += extra_rows

        if max_row >= insert_at:
            max_row += extra_rows

        sheet.merge_cells(
            start_row=min_row,
            start_column=min_col,
            end_row=max_row,
            end_column=max_col,
        )

    # Create rows matching the original item-row layout.
    for row in range(insert_at, insert_at + extra_rows):
        copy_row_format(sheet, LAST_ITEM_ROW, row)

        # Description, unit price and line total use merged cells.
        sheet.merge_cells(
            start_row=row, start_column=3,
            end_row=row, end_column=8
        )
        sheet.merge_cells(
            start_row=row, start_column=11,
            end_row=row, end_column=12
        )
        sheet.merge_cells(
            start_row=row, start_column=13,
            end_row=row, end_column=14
        )


def set_cell_value(sheet, row, column, value):
    """Write only to the top-left cell of a merged range."""
    cell = sheet.cell(row, column)

    if isinstance(cell, MergedCell):
        for merged_range in sheet.merged_cells.ranges:
            if (
                merged_range.min_row <= row <= merged_range.max_row
                and merged_range.min_col <= column <= merged_range.max_col
            ):
                cell = sheet.cell(
                    merged_range.min_row,
                    merged_range.min_col,
                )
                break

    cell.value = value


def find_label_row(sheet, keywords):
    """Find a totals row by searching its visible labels."""
    for row in range(1, sheet.max_row + 1):
        row_text = " ".join(
            str(sheet.cell(row, column).value or "").lower()
            for column in range(1, 13)
        )

        if any(keyword in row_text for keyword in keywords):
            return row

    return None


def update_totals(sheet, first_item_row, last_item_row):
    """Update totals formulas when the corresponding labels are found."""
    subtotal_row = find_label_row(sheet, ["subtotal", "sub total"])
    tax_row = find_label_row(sheet, ["gst", "vat"])
    grand_total_row = find_label_row(
        sheet, ["grand total", "total payable", "net total"]
    )

    if subtotal_row is not None:
        set_cell_value(
            sheet,
            subtotal_row,
            COL_LINE_TOTAL,
            f"=SUM(M{first_item_row}:M{last_item_row})",
        )

    if tax_row is not None and subtotal_row is not None:
        # Retain the template's existing 18% tax assumption for review.
        set_cell_value(
            sheet,
            tax_row,
            COL_LINE_TOTAL,
            f"=M{subtotal_row}*18%",
        )

    if grand_total_row is not None and subtotal_row is not None:
        if tax_row is not None:
            formula = f"=M{subtotal_row}+M{tax_row}"
        else:
            formula = f"=M{subtotal_row}"

        set_cell_value(
            sheet,
            grand_total_row,
            COL_LINE_TOTAL,
            formula,
        )


def generate_quotation(rfq_path=RFQ_PATH, output_path=None):
    """Create a draft Excel quotation from RFQ and supplier price data."""
    rfq_path = Path(rfq_path)

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Quotation template not found: {TEMPLATE_PATH}"
        )

    if not rfq_path.exists():
        raise FileNotFoundError(f"RFQ not found: {rfq_path}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    if output_path is None:
        output_path = OUTPUT_DIR / "IEPCI_Quotation_DRAFT.xlsx"
    else:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

    # 1. Extract all RFQ items.
    rfq_df = extract_rfq_lines(rfq_path)

    if rfq_df is None or rfq_df.empty:
        raise ValueError("No RFQ items were extracted.")

    # 2. Load supplier prices and match them to the RFQ.
    catalogue = load_catalogue()
    priced_df = add_supplier_prices(rfq_df, catalogue)

    # 3. Open the original template.
    workbook = load_workbook(TEMPLATE_PATH)

    try:
        if "Format" not in workbook.sheetnames:
            raise ValueError(
                "The quotation template has no 'Format' worksheet."
            )

        sheet = workbook["Format"]

        # 4. Make room for every RFQ item.
        expand_item_table(sheet, len(priced_df))

        last_item_row = FIRST_ITEM_ROW + len(priced_df) - 1

        # 5. Mark the output clearly as a draft.
        sheet["A5"] = (
            "QUOTATION - DRAFT - VERIFY PRICES, CURRENCY AND TAX"
        )
        sheet["M6"] = "IEPCI/QTN/2026/DRAFT-01"
        sheet["M8"] = "KBS/RFQ/2026/0147"

        # 6. Populate the item table.
        for index, (_, item) in enumerate(priced_df.iterrows()):
            row = FIRST_ITEM_ROW + index

            part_code = get_value(item, "part_code", default="")
            description = get_value(item, "description", default="")
            make = get_value(item, "make", default="")
            unit = get_value(item, "unit", default="Each")
            quantity = get_value(
                item, "qty", "quantity", default=None
            )

            if make and description:
                description = f"{make} - {description}"

            set_cell_value(sheet, row, COL_NUMBER, index + 1)
            set_cell_value(sheet, row, COL_PART_CODE, part_code)
            set_cell_value(sheet, row, COL_DESCRIPTION, description)
            set_cell_value(sheet, row, COL_UNIT, unit)
            set_cell_value(sheet, row, COL_QUANTITY, quantity)

            status = str(
                get_value(
                    item, "supplier_price_status", default=""
                )
            ).strip()

            price = get_value(item, "supplier_price_inr")

            # Only populate prices for the explicitly recognised
            # matching statuses. Missing/incomplete codes stay blank.
            approved_statuses = {
                "LIST PRICE FOUND - verify specs",
                "DUPLICATE CODE - same price",
            }

            if (
                status in approved_statuses
                and price is not None
                and pd.notna(price)
            ):
                set_cell_value(
                    sheet, row, COL_UNIT_PRICE, float(price)
                )
                set_cell_value(
                    sheet,
                    row,
                    COL_LINE_TOTAL,
                    f"=J{row}*K{row}",
                )
            else:
                set_cell_value(sheet, row, COL_UNIT_PRICE, None)
                set_cell_value(sheet, row, COL_LINE_TOTAL, None)

        # 7. Refresh totals formulas where template labels are recognised.
        update_totals(sheet, FIRST_ITEM_ROW, last_item_row)

        # Ask Excel to recalculate formulas when the workbook opens.
        try:
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.calculation.calcMode = "auto"
        except AttributeError:
            pass

        workbook.save(output_path)

    finally:
        workbook.close()

    print(f"Draft quotation created: {output_path}")
    print(f"RFQ items processed: {len(priced_df)}")
    print("Review unmatched codes and blank prices before proceeding.")
    print("Verify GBP/INR currency, tax and commercial terms before sending.")

    return output_path


if __name__ == "__main__":
    generate_quotation()
