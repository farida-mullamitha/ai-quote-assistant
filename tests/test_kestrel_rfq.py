"""Golden test for the Kestrel RFQ (KBS/RFQ/2026/0147).

The expected values below were checked by hand against the ABB (Feb 2026) and
Schneider (Apr 2026) price lists. If a test fails, either the code broke or
the price lists / RFQ changed: look at which before "fixing" the test.

Run from the project root:
    python -m pytest
"""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.check_rfq_codes import check_lines
from src.extract_rfq import extract_rfq_header, extract_rfq_lines
from src.generate_quotation import generate_quotation
from src.load_price_lists import load_catalogue
from src.match_missing_codes import find_missing_code_candidates

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
RFQ = RAW / "RFQ_KBS-0147.pdf"
ABB = RAW / "ABB_Price_List_Feb2026.xlsx"
SCHNEIDER = RAW / "Schneider LV Price list Apr 2026.xlsx"
TEMPLATE = RAW / "IEPCI Quotation .xlsx"

pytestmark = pytest.mark.skipif(
    not all(p.exists() for p in (RFQ, ABB, SCHNEIDER, TEMPLATE)),
    reason="Put the RFQ, both price lists and the template in data/raw/ first.",
)

# RFQ line -> list price in INR, as shown by the coverage check.
EXPECTED_PRICES = {
    "1": "2,280", "2": "440", "3": "6,280", "4": "19,900", "5": "7,840",
    "6": "3,540", "7": "22,710", "8": "8,150", "9": "2,680", "10": "290",
    "11": "14,530", "13": "19,780", "14": "59,350", "15": "4,120",
    "16": "1,430", "17": "32,370", "19": "45,240", "20": "23,070",
}


@pytest.fixture(scope="module")
def lines():
    return extract_rfq_lines(RFQ)


@pytest.fixture(scope="module")
def header():
    return extract_rfq_header(RFQ)


@pytest.fixture(scope="module")
def catalogue():
    return load_catalogue(ABB, SCHNEIDER)


@pytest.fixture(scope="module")
def report(lines, catalogue):
    return check_lines(lines, catalogue)


# ------------------------------------------------------------ extraction
def test_extracts_all_twenty_lines(lines):
    assert list(lines["line_no"]) == [str(n) for n in range(1, 21)]
    assert set(lines["make"]) == {"ABB", "Schneider Electric"}
    assert lines.loc[lines["line_no"] == "1", "qty"].item() == 24


def test_part_code_cells_are_split_into_code_and_note(lines):
    line5 = lines[lines["line_no"] == "5"].iloc[0]
    assert line5["part_code"] == "1SBL237001RXX00"
    assert line5["code_note"] == "coil 230V AC"

    for number in ("12", "18"):
        line = lines[lines["line_no"] == number].iloc[0]
        assert line["part_code"] == ""
        assert "confirmed" in line["code_note"]


def test_header_details(header):
    assert header["customer_name"] == "Kestrel Building Services Ltd"
    assert header["rfq_number"] == "KBS/RFQ/2026/0147"
    assert header["rfq_date"] == "15/10/2026"
    assert header["currency"] == "GBP"
    assert header["email"] == "procurement@kestrel-bs.example"


# ------------------------------------------------------------ catalogue check
def test_catalogue_loaded_both_suppliers(catalogue):
    assert set(catalogue["supplier"]) == {"ABB", "Schneider Electric"}
    assert len(catalogue) > 6000


@pytest.mark.parametrize("line_no, price", EXPECTED_PRICES.items())
def test_list_price_found_for_each_coded_line(report, line_no, price):
    first_row = report[report["line"] == line_no].iloc[0]
    assert first_row["price_inr"] == price


def test_lines_that_need_a_person(report):
    def status(line_no):
        return report[report["line"] == line_no].iloc[0]["result"]

    assert status("4").startswith("MATCHED")
    assert "INCOMPLETE CODE" in status("5")
    assert "MISSING CODE" in status("12")
    assert "MISSING CODE" in status("18")
    assert "DUPLICATE MATCH" in status("15")   # listed under CVS and NSX
    assert "DUPLICATE MATCH" in status("16")


# ------------------------------------------------------------ missing codes
def top_candidate(candidates, line_no):
    return next(c for c in candidates if c["line_no"] == line_no and c["rank"] == 1)


def test_missing_code_candidates_rank_the_right_product_first(lines, catalogue):
    candidates = find_missing_code_candidates(lines, catalogue)

    # Line 18: electrically operated EasyPact NW 1600 A, TS6S (LSIG), 50 kA.
    # The manual-operation twin (ENW216N3MW6S) must NOT rank first.
    line18 = top_candidate(candidates, "18")
    assert line18["part_code"] == "ENW216N3NW6S"
    assert line18["price"] == 581270

    # Line 12: ABB Formula Air FA1 EK1 LSIG 1000 A draw-out (ABB lists no part
    # code for it, so identify it by price). EK2 (352,400) must not win.
    assert top_candidate(candidates, "12")["price"] == 338400


# ------------------------------------------------------------ quotation
def test_quotation_has_header_and_a_total_for_every_line(lines, header, tmp_path):
    approved = lines.assign(unit_rate=100.0, price_approved=True)
    output = generate_quotation(
        output_path=tmp_path / "quote.xlsx", rfq_df=approved, header=header
    )
    ws = load_workbook(output)["Format"]

    first, last = 13, 13 + len(lines) - 1                  # rows 13..32
    for row in range(first, last + 1):
        # Regression test: rows beyond the template's 17 once had no total.
        assert str(ws.cell(row, 13).value).startswith("="), f"row {row}"
        assert ws.cell(row, 11).value == 100.0

    assert ws.cell(last + 1, 13).value == (
        f'=IF(COUNT(K{first}:K{last})=0,"",SUM(M{first}:M{last}))'
    )
    assert ws["C6"].value == "Kestrel Building Services Ltd"
    assert ws["M8"].value == "KBS/RFQ/2026/0147"
    assert "GBP" in ws["A5"].value                          # currency warning


def test_unapproved_rates_are_never_filled_in(lines, header, tmp_path):
    unapproved = lines.assign(unit_rate=999.0, price_approved=False)
    output = generate_quotation(
        output_path=tmp_path / "quote.xlsx", rfq_df=unapproved, header=header
    )
    ws = load_workbook(output)["Format"]
    assert all(ws.cell(row, 11).value is None for row in range(13, 33))
